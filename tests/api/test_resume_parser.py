"""Labeled parser checks and real document extraction, without a database."""

from contextlib import redirect_stderr, redirect_stdout
from io import BytesIO, StringIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
import typst

from api.resume.parser import (
    MAX_FILE_BYTES, ParseError, ParsedResume, main, parse_document, parse_text,
)

FIXTURES = Path(__file__).parents[1] / "fixtures" / "resume_parser.json"


def field_value(result, path):
    value = result
    for part in path.split("."):
        try:
            value = value[int(part)] if isinstance(value, list) else value[part]
        except (KeyError, IndexError, TypeError):
            return None
    return value


def pdf(source):
    return typst.compile((
        '#set page(paper: "us-letter", margin: 40pt)\n'
        '#set text(size: 11pt)\n#set par(leading: 5pt)\n' + source
    ).encode(), root="/")


def docx(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def add_link(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    relationship = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


class LabeledFixtureTests(unittest.TestCase):
    def test_labeled_text_fixtures(self):
        incorrect, missed, checked = [], [], 0
        for fixture in json.loads(FIXTURES.read_text()):
            with self.subTest(fixture=fixture["name"]):
                parsed = parse_text(fixture["text"])
                result = parsed.model_dump()
                self.assertEqual(result, parse_text(fixture["text"]).model_dump())
                self.assertEqual(ParsedResume.model_validate_json(parsed.model_dump_json()), parsed)
                for path, expected in fixture["expected"].items():
                    checked += 1
                    actual = field_value(result, path)
                    if actual != expected:
                        issue = (fixture["name"], path, expected, actual)
                        (missed if expected not in ("", None) and actual in ("", None) else incorrect).append(issue)
                self.assertEqual(len(parsed.entries), fixture["entry_count"])
                self.assertEqual(len(parsed.skills), fixture["skill_count"])
                self.assertEqual(set(w.code for w in parsed.warnings), set(fixture["warnings"]))
                text = "\n".join(item.text for item in parsed.unparsed)
                for expected in fixture.get("unparsed_contains", []):
                    self.assertIn(expected, text)
                source_ids = {block.id for block in parsed.sources}
                self.assertTrue(all(set(item.source_ids) <= source_ids for item in parsed.unparsed))
                self.assertTrue(all(set(item.source_ids) <= source_ids for item in parsed.warnings))
                self.assertTrue(all(
                    field_value(result, path) not in ("", None)
                    for block in parsed.sources for path in block.fields if ":" not in path
                ))
        print(f"\nLabeled text fixtures: {checked} checked fields; {len(incorrect)} incorrect; {len(missed)} missed.")
        self.assertEqual(incorrect, [], f"Incorrect fields: {incorrect}")
        self.assertEqual(missed, [], f"Missed fields: {missed}")

    def test_no_content_is_silently_lost(self):
        result = parse_text("Alex Morgan\n\nUnusual Section\nSomething not understood.\n")
        self.assertIn("Something not understood.", "\n".join(item.text for item in result.unparsed))
        self.assertEqual([block.text for block in result.sources],
                         ["Alex Morgan", "Unusual Section", "Something not understood."])

    def test_long_values_are_preserved_and_flagged(self):
        bullet = "Built " + "a" * 650 + "."
        tool = "Custom" + "x" * 130
        result = parse_text(
            f"Alex Morgan\nExperience\nEngineer | Acme | 2020 – Present\n- {bullet}\nSkills\nTools: {tool}\n")
        self.assertEqual(result.entries[0].bullets[0].text, bullet)
        self.assertEqual(result.skills[0].name, tool)
        self.assertEqual({w.field for w in result.warnings if w.code == "field_too_long"},
                         {"entries.0.bullets.0.text", "skills.0.name"})
        self.assertTrue(all(w.source_ids for w in result.warnings))

    def test_orphan_bullet_is_not_a_fabricated_entry(self):
        result = parse_text("Experience\n- Built something.\n")
        self.assertEqual(result.entries, [])
        self.assertIn("orphan_bullet", {w.code for w in result.warnings})
        self.assertIn("Built something.", result.unparsed[0].text)

    def test_empty_text_is_valid_but_limited(self):
        result = parse_text("")
        self.assertEqual(result.sources, [])
        self.assertEqual([w.code for w in result.warnings], ["limited_extraction"])

    def test_duplicate_skills_retain_all_source_references(self):
        result = parse_text("Skills\nLanguages: Python, Python\nLanguages: Python\n")
        self.assertEqual(len(result.skills), 1)
        self.assertEqual(len([block for block in result.sources if "skills.0.name" in block.fields]), 2)

    def test_ambiguous_skill_alias_does_not_match_inside_a_word(self):
        result = parse_text("Skills\nDjangoish, Reacted, Custom Python Service\n")
        self.assertEqual([s.category for s in result.skills], [None, None, None])

    def test_multiple_conflicting_organizations_are_unset(self):
        result = parse_text(
            "Experience\nEngineer | Acme Inc | Widgets LLC | 2024 – Present\n- Built something.\n")
        self.assertEqual(result.entries[0].org, "")
        self.assertIn("ambiguous_field", {w.code for w in result.warnings})

    def test_inline_skill_groups_and_spaced_slashes(self):
        result = parse_text("Skills\nLanguages: Python / SQL; Tools: Git, CI/CD\n")
        self.assertEqual([(s.name, s.category) for s in result.skills],
                         [("Python", "language"), ("SQL", "language"), ("Git", "tool"), ("CI/CD", "tool")])

    def test_comma_degree_dash_pair_and_organization_location(self):
        result = parse_text(
            "Education\nUniversity of Washington, B.S. Computer Science | 2023 – 2027\n"
            "Experience\nSoftware Engineer — Northstar Labs, Seattle, WA | 2024 – Present\n- Built an API.\n")
        self.assertEqual(result.entries[0].org, "University of Washington")
        self.assertEqual(result.entries[0].role, "B.S. Computer Science")
        self.assertEqual(result.entries[1].org, "Northstar Labs")
        self.assertEqual(result.entries[1].role, "Software Engineer")
        self.assertEqual(result.entries[1].location, "Seattle, WA")

    def test_discarded_headers_do_not_point_to_later_entries(self):
        result = parse_text("Education\nGPA: 3.5\nExperience\nEngineer | Acme | 2024 – Present\n- Built an API.\n")
        for source in result.sources:
            if source.text == "GPA: 3.5":
                self.assertNotIn("entries.0.detail", source.fields)

    def test_three_roles_share_only_explicit_parent_heading(self):
        result = parse_text(
            "Experience\nNorthstar Labs\n  Senior Engineer | 2024 – Present\n  - Led a team.\n"
            "  Engineer | 2022 – 2024\n  - Built an API.\n"
            "  Intern | 2021 – 2022\n  - Built tests.\n")
        self.assertEqual([e.org for e in result.entries], ["Northstar Labs"] * 3)

    def test_unseparated_mixed_headers_are_not_guessed(self):
        result = parse_text(
            "Education\nUniversity of Washington B.S. Computer Science | 2023 – 2027\n"
            "Experience\nSoftware Engineer Northstar Labs | 2024 – Present\n- Built an API.\n")
        self.assertEqual([(e.org, e.role) for e in result.entries], [("", ""), ("", "")])
        self.assertIn("ambiguous_entry_header", {w.code for w in result.warnings})
        self.assertIn("University of Washington B.S. Computer Science", "\n".join(u.text for u in result.unparsed))

    def test_unknown_category_does_not_hide_overlength_skill(self):
        name = "x" * 150
        result = parse_text("Skills\n" + name)
        self.assertEqual(result.skills[0].name, name)
        self.assertEqual({w.code for w in result.warnings},
                         {"field_too_long", "unresolved_skill_category"})


class PdfExtractionTests(unittest.TestCase):
    def test_real_pdf_fields_dates_links_and_bullet_wraps(self):
        data = pdf(
            "*Alex Morgan*\\\nalex\\@example.com | Seattle, WA\\\n"
            '#link("https://github.com/alex")[GitHub]\n\n'
            "== Education\nUniversity of Washington | B.S. Computer Science\\\n"
            "Sep 2023 – Jun 2027\\\nGPA: 3.85 / 4.0\n\n"
            "== Experience\nSoftware Engineer | Acme | 2024 – Present\n\n"
            "- Built a cache that processed 12,000 requests per second and reduced "
            "tail latency by 35% across all production regions while preserving correctness.\n"
        )
        result = parse_document(data, "pdf")
        self.assertEqual(result.profile.full_name, "Alex Morgan")
        self.assertEqual(result.profile.email, "alex@example.com")
        self.assertEqual(result.profile.links[0].url, "https://github.com/alex")
        self.assertEqual(result.entries[0].org, "University of Washington")
        self.assertEqual(result.entries[0].role, "B.S. Computer Science")
        self.assertEqual(result.entries[0].detail, "3.85 / 4.0")
        self.assertEqual(result.entries[1].org, "Acme")
        self.assertEqual(result.entries[1].bullets[0].text,
                         "Built a cache that processed 12,000 requests per second and reduced "
                         "tail latency by 35% across all production regions while preserving correctness.")
        self.assertTrue(all(b.page == 1 and b.bbox is not None for b in result.sources))
        self.assertEqual(result.model_dump(), parse_document(data, "pdf").model_dump())

    def test_two_columns_are_separate_streams(self):
        data = pdf(
            "*Alex Morgan*\\\nalex\\@example.com\n\n"
            "#grid(columns: (1fr, 1fr), gutter: 50pt,\n"
            "[== Education\nUniversity of Washington\\\nB.S. Computer Science\\\n2023 – 2027\n\n"
            "== Skills\nLanguages: Python, C++\\\nTools: Git, Docker],\n"
            "[== Experience\nEngineer | Acme\\\n2024 – Present\n\n"
            "- Built a service.\n- Reduced costs by 20%.\n\n"
            "== Projects\nTrailhead\\\n2025 – Present\n\n- Built a map cache.])"
        )
        result = parse_document(data, "pdf")
        self.assertEqual([(e.kind, e.org) for e in result.entries],
                         [("education", "University of Washington"), ("experience", "Acme"), ("project", "Trailhead")])
        self.assertEqual([s.name for s in result.skills], ["Python", "C++", "Git", "Docker"])
        self.assertEqual([b.text for b in result.entries[1].bullets],
                         ["Built a service.", "Reduced costs by 20%."])
        self.assertEqual(result.entries[2].bullets[0].text, "Built a map cache.")
        self.assertEqual({b.column for b in result.sources}, {-1, 0, 1})

    def test_right_aligned_dates_are_not_a_second_column(self):
        data = pdf(
            "Alex Morgan\n\n== Experience\n"
            '#grid(columns: (1fr, auto), [Engineer | Acme], [2024 – Present])\n'
            "- Built a service.\n\n"
            '#grid(columns: (1fr, auto), [Developer | Widgets], [2022 – 2024])\n'
            "- Built another service.\n"
        )
        result = parse_document(data, "pdf")
        self.assertEqual([(e.org, e.start_date, e.end_date) for e in result.entries],
                         [("Acme", "2024", "Present"), ("Widgets", "2022", "2024")])
        self.assertNotIn(1, {b.column for b in result.sources})

    def test_repeated_headers_and_page_numbers_preserve_continued_bullets(self):
        data = pdf(
            '#set page(header: [Alex Morgan], footer: context counter(page).display())\n'
            "== Experience\nEngineer | Acme | 2024 – Present\n\n- Built a service.\n"
            "#pagebreak()\n- Reduced costs by 20%.\n"
        )
        result = parse_document(data, "pdf")
        self.assertEqual(len(result.entries), 1)
        self.assertEqual([b.text for b in result.entries[0].bullets],
                         ["Built a service.", "Reduced costs by 20%."])
        self.assertIn("repeated_header_footer", {b.ignored for b in result.sources})
        self.assertIn("page_number", {b.ignored for b in result.sources})

    def test_mixed_image_and_text_document_is_flagged(self):
        from PIL import Image
        stream = BytesIO()
        Image.new("RGB", (20, 20), "white").save(stream, format="PNG")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.png"
            path.write_bytes(stream.getvalue())
            data = pdf("Alex Morgan\n#pagebreak()\n" + f'#image("{path.as_posix()}")')
        result = parse_document(data, "pdf")
        self.assertEqual(result.profile.full_name, "Alex Morgan")
        self.assertIn("unsupported_layout", {w.code for w in result.warnings})
        self.assertEqual(result.unparsed[-1].reason, "unsupported_layout")


class DocxExtractionTests(unittest.TestCase):
    def test_tables_headers_hyperlinks_and_numbered_bullets(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Alex Morgan"
        document.add_paragraph("alex@example.com | Seattle, WA")
        add_link(document.add_paragraph(), "Portfolio", "https://alex.dev")
        document.add_heading("Education", 1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "University of Washington\nB.S. Computer Science"
        table.cell(0, 1).text = "Sep 2023 – Jun 2027\nGPA: 3.85"
        document.add_heading("Experience", 1)
        document.add_paragraph("Engineer | Acme | 2024 – Present")
        document.add_paragraph("Built an API.", style="List Bullet")
        document.add_paragraph("Reduced costs by 25%.", style="List Number")
        result = parse_document(docx(document), "docx")
        self.assertEqual(result.profile.full_name, "Alex Morgan")
        self.assertEqual(result.profile.links[0].url, "https://alex.dev")
        self.assertEqual(result.entries[0].org, "University of Washington")
        self.assertEqual(result.entries[0].end_date, "Jun 2027")
        self.assertEqual(result.entries[0].detail, "3.85")
        self.assertEqual(result.entries[1].org, "Acme")
        self.assertEqual([b.text for b in result.entries[1].bullets],
                         ["Built an API.", "Reduced costs by 25%."])
        self.assertTrue(any("/cell:" in b.container for b in result.sources))
        self.assertEqual(result.model_dump(), parse_document(docx(document), "docx").model_dump())

    def test_two_column_table_and_nested_cells(self):
        document = Document()
        document.add_paragraph("Alex Morgan")
        layout = document.add_table(rows=1, cols=2)
        left, right = layout.rows[0].cells
        left.paragraphs[0].text = "Education"
        left.add_paragraph("University of Washington")
        left.add_paragraph("B.S. Computer Science")
        left.add_paragraph("2023 – 2027")
        left.add_paragraph("Skills")
        nested = left.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Languages: Python, SQL"
        right.paragraphs[0].text = "Experience"
        right.add_paragraph("Engineer | Acme | 2024 – Present")
        right.add_paragraph("Built a service.", style="List Bullet")
        result = parse_document(docx(document), "docx")
        self.assertEqual([(e.kind, e.org) for e in result.entries],
                         [("education", "University of Washington"), ("experience", "Acme")])
        self.assertEqual([s.name for s in result.skills], ["Python", "SQL"])
        self.assertEqual(result.entries[1].bullets[0].text, "Built a service.")

    def test_multiple_roles_under_bold_employer(self):
        document = Document()
        document.add_heading("Experience", 1)
        document.add_paragraph().add_run("Northstar Labs").bold = True
        document.add_paragraph("Senior Engineer | 2024 – Present")
        document.add_paragraph("Led a migration.", style="List Bullet")
        document.add_paragraph("Engineer | 2022 – 2024")
        document.add_paragraph("Built an API.", style="List Bullet")
        result = parse_document(docx(document), "docx")
        self.assertEqual([e.org for e in result.entries], ["Northstar Labs", "Northstar Labs"])
        self.assertEqual([e.role for e in result.entries], ["Senior Engineer", "Engineer"])

    def test_textbox_is_preserved_for_review(self):
        document = Document()
        document.add_paragraph("Alex Morgan")
        paragraph = document.add_paragraph()
        box = OxmlElement("w:txbxContent")
        inner = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Awards: First place, 2025"
        run.append(text)
        inner.append(run)
        box.append(inner)
        paragraph._p.append(box)
        result = parse_document(docx(document), "docx")
        self.assertEqual(result.entries, [])
        self.assertIn("Awards: First place, 2025", "\n".join(b.text for b in result.unparsed))
        self.assertIn("unsupported_layout", {w.code for w in result.warnings})

    def test_merged_cells_are_not_duplicated(self):
        document = Document()
        document.add_heading("Skills", 1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Tools: Git, Docker"
        result = parse_document(docx(document), "docx")
        self.assertEqual([s.name for s in result.skills], ["Git", "Docker"])
        self.assertEqual(len([b for b in result.sources if b.text == "Tools: Git, Docker"]), 1)

    def test_bullet_paragraph_continuation(self):
        document = Document()
        document.add_heading("Experience", 1)
        document.add_paragraph("Engineer | Acme | 2024 – Present")
        document.add_paragraph("Built a cache for", style="List Bullet")
        continuation = document.add_paragraph("12,000 requests per second.")
        continuation.paragraph_format.left_indent = Inches(.5)
        result = parse_document(docx(document), "docx")
        self.assertEqual(result.entries[0].bullets[0].text, "Built a cache for 12,000 requests per second.")

    def test_mailto_link_and_footer_are_not_skills(self):
        document = Document()
        document.add_paragraph("Alex Morgan")
        add_link(document.add_paragraph(), "Email", "mailto:alex@example.com")
        document.add_heading("Skills", 1)
        document.add_paragraph("Languages: Python")
        document.sections[0].footer.paragraphs[0].text = "Alex Morgan"
        result = parse_document(docx(document), "docx")
        self.assertEqual(result.profile.email, "alex@example.com")
        self.assertEqual([s.name for s in result.skills], ["Python"])
        self.assertIn("Alex Morgan", "\n".join(b.text for b in result.unparsed))


class InputAndCliTests(unittest.TestCase):
    def assert_error(self, code, callback):
        with self.assertRaises(ParseError) as caught:
            callback()
        self.assertEqual(caught.exception.code, code)

    def test_malformed_and_unsupported_inputs(self):
        self.assert_error("malformed", lambda: parse_document(b"not a pdf", "pdf"))
        self.assert_error("malformed", lambda: parse_document(b"not a docx", "docx"))
        self.assert_error("unsupported_type", lambda: parse_document(b"text", "txt"))
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("wrong.xml", "<document />")
        self.assert_error("malformed", lambda: parse_document(stream.getvalue(), "docx"))
        source = docx(Document())
        broken = BytesIO()
        with ZipFile(BytesIO(source)) as archive, ZipFile(broken, "w") as target:
            for item in archive.infolist():
                target.writestr(item.filename, b"<broken" if item.filename == "word/document.xml" else archive.read(item.filename))
        self.assert_error("malformed", lambda: parse_document(broken.getvalue(), "docx"))

    def test_oversized_inputs(self):
        data = b"x" * (MAX_FILE_BYTES + 1)
        self.assert_error("too_large", lambda: parse_document(data, "pdf"))
        self.assert_error("too_large", lambda: parse_text(data.decode()))

    def test_encrypted_pdf_and_legacy_word(self):
        # Construct an actual Standard Security Handler dictionary; pdfminer
        # detects the password failure before attempting content extraction.
        data = b"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [] /Count 0 >> endobj
3 0 obj << /Filter /Standard /V 1 /R 2 /O (12345678901234567890123456789012) /U (12345678901234567890123456789012) /P -4 >> endobj
trailer << /Root 1 0 R /Encrypt 3 0 R /ID [(0123456789abcdef)(0123456789abcdef)] >>
%%EOF
"""
        self.assert_error("encrypted", lambda: parse_document(data, "pdf"))
        self.assert_error("encrypted", lambda: parse_document(bytes.fromhex("d0cf11e0a1b11ae1"), "docx"))

    def test_image_only_pdf_and_docx(self):
        from PIL import Image
        stream = BytesIO()
        Image.new("RGB", (20, 20), "white").save(stream, format="PNG")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.png"
            path.write_bytes(stream.getvalue())
            self.assert_error("image_only", lambda: parse_document(pdf(f'#image("{path.as_posix()}")'), "pdf"))
        document = Document()
        document.add_picture(BytesIO(stream.getvalue()))
        self.assert_error("image_only", lambda: parse_document(docx(document), "docx"))

    def test_empty_readable_documents_have_limited_extraction(self):
        for data, kind in ((pdf(""), "pdf"), (docx(Document()), "docx")):
            with self.subTest(kind=kind):
                result = parse_document(data, kind)
                self.assertIn("limited_extraction", {w.code for w in result.warnings})

    def test_cli_json_and_error_exit_codes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.txt"
            path.write_text("Alex Morgan\nalex@example.com\n", encoding="utf-8")
            process = subprocess.run(
                [sys.executable, "-m", "api.resume.parser", str(path)],
                capture_output=True, text=True, check=False)
            self.assertEqual(process.returncode, 0, process.stderr)
            self.assertEqual(json.loads(process.stdout)["profile"]["email"], "alex@example.com")
            self.assertEqual(process.stderr, "")
            path.write_bytes(b"\xff")
            stdout, stderr = StringIO(), StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                status = main([str(path)])
            self.assertEqual(status, 1)
            self.assertEqual(stdout.getvalue(), "")
            self.assertNotIn("Traceback", stderr.getvalue())
            with redirect_stderr(stderr):
                self.assertEqual(main([str(Path(directory) / "missing.txt")]), 1)
                self.assertEqual(main([str(Path(directory) / "resume.exe")]), 1)


if __name__ == "__main__":
    unittest.main()
