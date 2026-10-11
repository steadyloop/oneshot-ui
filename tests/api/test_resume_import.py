"""Public parsing must remain stateless and bound streamed uploads."""
from io import BytesIO
import unittest

from docx import Document
from api.resume.render import compile_pdf
from unittest.mock import patch

import httpx
from fastapi import FastAPI

from api.resume.import_routes import router


class ResumeImportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        app = FastAPI()
        app.include_router(router)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()

    async def test_guest_text_returns_reviewable_draft(self):
        response = await self.client.post('/api/resume/parse?file_type=text', content='Alex Morgan\nalex@example.com\nSkills\nLanguages: Python')
        self.assertEqual(response.status_code, 200)
        draft = response.json()
        self.assertEqual(draft['profile']['email'], 'alex@example.com')
        self.assertTrue(draft['sources'])
        self.assertEqual(draft['skills'][0]['name'], 'Python')

    async def test_malformed_and_invalid_unicode(self):
        for kind, body, code in [('pdf', b'broken', 'malformed'), ('docx', b'broken', 'malformed'), ('text', b'\xff', 'invalid_text')]:
            response = await self.client.post(f'/api/resume/parse?file_type={kind}', content=body)
            self.assertEqual(response.status_code, 422)
            self.assertEqual(response.json()['error']['code'], code)

    async def test_explicit_type_required(self):
        for query in ['', '?file_type=html']:
            response = await self.client.post('/api/resume/parse' + query, content=b'x')
            self.assertEqual(response.status_code, 422)

    async def test_content_length_limit(self):
        response = await self.client.post('/api/resume/parse?file_type=text', content=b'x', headers={'content-length': str(11 * 1024 * 1024)})
        self.assertEqual(response.status_code, 413)

    async def test_chunked_upload_limit(self):
        async def chunks():
            yield b'1234'
            yield b'5678'
        with patch('api.resume.import_routes.MAX_FILE_BYTES', 6):
            response = await self.client.post('/api/resume/parse?file_type=text', content=chunks())
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()['error']['code'], 'too_large')

    async def test_real_documents(self):
        document = Document()
        document.add_paragraph('Alex Morgan')
        document.add_paragraph('alex@example.com')
        buffer = BytesIO()
        document.save(buffer)
        pdf, _, error = compile_pdf('Alex Morgan\n\n#link("mailto:alex@example.com")[alex\\@example.com]')
        self.assertFalse(error)
        for kind, data in [('docx', buffer.getvalue()), ('pdf', pdf)]:
            response = await self.client.post(f'/api/resume/parse?file_type={kind}', content=data)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['profile']['email'], 'alex@example.com')
